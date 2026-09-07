"""
crear_parametros.py
-------------------
Genera 'Parametros_Modelo.xlsx' con TODOS los datos del modelo, parametrizados y
explicados.  El modelo principal (modelo.py) lee la hoja 'Parametros'.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

AZUL = "1F4E78"; AZULC = "DDEBF7"; GRIS = "F2F2F2"; VERDE = "E2EFDA"
bold_w = Font(bold=True, color="FFFFFF")
titulo = Font(bold=True, size=13, color=AZUL)
fill_h = PatternFill("solid", fgColor=AZUL)
fill_s = PatternFill("solid", fgColor=AZULC)
fill_e = PatternFill("solid", fgColor=VERDE)
thin = Side(style="thin", color="BFBFBF")
borde = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center")


def _hdr(ws, row, cols):
    for j, t in enumerate(cols, 1):
        c = ws.cell(row, j, t); c.font = bold_w; c.fill = fill_h
        c.alignment = center; c.border = borde


def _row(ws, row, vals, fill=None):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row, j, v); c.border = borde; c.alignment = wrap
        if fill:
            c.fill = fill


# ---------------------------------------------------------------- PARAMETROS
PARAMS = [
    # clave, valor, unidad, descripcion
    ("--- RUTAS DE DATOS ---", "", "", ""),
    ("ruta_afluentes_ralco",  "datos/S0130134-25_26.xlsx", "texto",
     "Ruta al Excel de caudales afluentes al embalse Ralco (caudales medios semanales, m3/s)."),
    ("ruta_afluentes_pangue", "datos/S0130514-25_26.xlsx", "texto",
     "Ruta al Excel de caudales de la hoya intermedia Pangue-Ralco (C.I. Bio Bio, m3/s)."),
    ("carpeta_salidas", "salidas", "texto", "Carpeta donde se guardan los resultados."),

    ("--- CAUDAL ECOLOGICO ---", "", "", ""),
    ("caudal_ecologico_m3s", 27.1, "m3/s",
     "Caudal ecologico permanente que Ralco entrega y que turbina Palmucho; tambien llega a Pangue."),

    ("--- POTENCIAS ---", "", "", ""),
    ("ralco_pot_max_MW", 640.0, "MW", "Potencia maxima de Ralco."),
    ("ralco_pot_min_MW", 100.0, "MW", "Potencia minima de Ralco cuando opera (puede detenerse=0). [informativo v1]"),
    ("pangue_pot_max_MW", 400.0, "MW", "Potencia maxima de Pangue."),
    ("pangue_pot_min_MW", 20.0, "MW", "Potencia minima de Pangue. Pangue NO puede detenerse."),

    ("--- COTAS / VOLUMENES RALCO ---", "", "", ""),
    ("ralco_cota_ini", 725.0, "msnm", "Cota inicial de Ralco al 01-ene."),
    ("ralco_cota_min_op", 692.0, "msnm", "Cota minima de operacion de Ralco."),
    ("ralco_cota_max", 725.0, "msnm", "Cota maxima de Ralco."),
    ("ralco_cota_umbral_baja", 710.0, "msnm", "Umbral: sobre esta cota la bajada diaria maxima es mayor."),
    ("ralco_baja_sobre_umbral_cm", 40.0, "cm/dia", "Bajada maxima de cota por dia SOBRE la cota umbral (710)."),
    ("ralco_baja_bajo_umbral_cm", 25.0, "cm/dia", "Bajada maxima de cota por dia BAJO la cota umbral (710)."),
    ("ralco_cota_vertimiento", 708.0, "msnm", "Ralco solo puede verter sobre esta cota. [restriccion v2/binaria]"),

    ("--- COTAS / VOLUMENES PANGUE ---", "", "", ""),
    ("pangue_cota_ini", 510.0, "msnm", "Cota inicial de Pangue al 01-ene."),
    ("pangue_cota_min_op", 507.0, "msnm", "Cota minima de operacion de Pangue."),
    ("pangue_cota_max", 510.0, "msnm", "Cota maxima de Pangue (opera sin restricciones de bajada)."),

    ("--- TEMPORADA DE RIEGO ---", "", "", ""),
    ("riego_mes_ini", 1, "mes", "Primer mes de riego (1=enero). En riego Ralco entrega al menos el afluente."),
    ("riego_mes_fin", 4, "mes", "Ultimo mes de riego (4=abril)."),

    ("--- MANTENIMIENTO RALCO ---", "", "", ""),
    ("mant_duraciones_dias", "120,150,180", "dias",
     "Duraciones a evaluar (base 120; sensibilidades 150 y 180). Ralco se detiene por completo."),
    ("mant_permite_cruce_anio", "SI", "SI/NO",
     "Si SI, la ventana de mantenimiento puede cruzar el fin de año (calendario ciclico)."),

    ("--- CONFIGURACION DE BUSQUEDA ---", "", "", ""),
    ("busqueda_paso_dias", 15, "dias", "Paso del barrido grueso de fechas de inicio del mantenimiento."),
    ("busqueda_refina_dias", 3, "dias", "Paso del refinamiento fino alrededor de la mejor fecha gruesa."),
    ("n_iter_busqueda", 1, "-", "Iteraciones de linealizacion en el barrido (1 = rapido)."),
    ("n_iter_final", 3, "-", "Iteraciones de linealizacion en el calculo final en la fecha optima."),
    ("n_bandas_probabilidad", 5, "-", "Nro de bandas de probabilidad (5 => tramos de 20%)."),
    ("criterio_bandas", "afluente", "afluente/vertimiento",
     "Variable para ordenar escenarios de humedo a seco. 'afluente'=caudal anual afluente."),
]

DIAS_SEMANA = [
    ("Largo del mes", "Dias sem1", "Dias sem2", "Dias sem3", "Dias sem4", "Total"),
    (31, 7, 8, 8, 8, 31),
    (30, 7, 8, 7, 8, 30),
    (28, 7, 7, 7, 7, 28),
    (29, 7, 7, 7, 8, 29),
]

MESES_CAL = [("ENE", 31), ("FEB", 28), ("MAR", 31), ("ABR", 30), ("MAY", 31), ("JUN", 30),
             ("JUL", 31), ("AGO", 31), ("SEP", 30), ("OCT", 31), ("NOV", 30), ("DIC", 31)]


def construir(path="Parametros_Modelo.xlsx"):
    wb = openpyxl.Workbook()

    # ---------- LEEME ----------
    ws = wb.active; ws.title = "LEEME"
    ws["A1"] = "MODELO DE OPTIMIZACION RALCO + PANGUE  (v1 - solo fisico)"
    ws["A1"].font = titulo
    texto = [
        "",
        "Objetivo: encontrar la FECHA OPTIMA de inicio del mantenimiento de Ralco que",
        "MINIMICE el vertimiento esperado (energia perdida, MWh) de Ralco + Pangue.",
        "",
        "Como usar:",
        "  1) Edite los valores de la hoja 'Parametros' (columna VALOR).",
        "  2) Copie los Excel de afluentes en la carpeta 'datos/' (o ajuste las rutas).",
        "  3) Ejecute:  python src/modelo.py",
        "  4) Revise los resultados en la carpeta 'salidas/'.",
        "",
        "Salidas principales:",
        "  - resumen_fecha_optima.csv     : fecha optima por duracion y KPIs.",
        "  - barrido_fechas.csv           : vertimiento esperado vs fecha de inicio (curva).",
        "  - bandas_probabilidad.csv      : resultados por banda de 20% (humedo->seco).",
        "  - cota_evolucion_bandas.csv    : evolucion diaria de cota por banda.",
        "  - detalle_diario_bandas.csv    : series diarias promedio por banda.",
        "",
        "Convenciones de unidades:  cota[msnm]  volumen[Mm3]  caudal[m3/s]  energia[MWh]",
        "Rendimientos usados: Rendimiento_Ralco2 y Rendimiento_Pangue2  [MW/(m3/s)].",
        "",
        "Bandas de probabilidad: los escenarios (años hidrologicos) se ordenan de HUMEDO",
        "a SECO. Numeros mas ALTOS = mas SECOS = menor vertimiento (80-100%).",
        "",
        "Nota: el modelo trabaja los embalses en VOLUMEN y solo convierte a cota para",
        "mostrar la equivalencia (evita el error de ida-y-vuelta cota<->volumen).",
    ]
    for i, t in enumerate(texto, 3):
        ws.cell(i, 1, t)
    ws.column_dimensions["A"].width = 95

    # ---------- PARAMETROS ----------
    ws = wb.create_sheet("Parametros")
    ws["A1"] = "PARAMETROS DEL MODELO"; ws["A1"].font = titulo
    _hdr(ws, 3, ["PARAMETRO", "VALOR", "UNIDAD", "DESCRIPCION"])
    r = 4
    for clave, val, uni, desc in PARAMS:
        if str(clave).startswith("---"):
            c = ws.cell(r, 1, clave.replace("-", "").strip())
            c.font = Font(bold=True, color=AZUL)
            for j in range(1, 5):
                ws.cell(r, j).fill = fill_s; ws.cell(r, j).border = borde
        else:
            _row(ws, r, [clave, val, uni, desc])
            ws.cell(r, 2).fill = fill_e   # celda editable en verde
            ws.cell(r, 2).alignment = center
        r += 1
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 75

    # ---------- CALENDARIO ----------
    ws = wb.create_sheet("Calendario")
    ws["A1"] = "CALENDARIO Y REPARTO SEMANA -> DIA"; ws["A1"].font = titulo
    ws["A3"] = "Reparto de dias por semana segun largo del mes (regla del usuario):"
    _hdr(ws, 4, list(DIAS_SEMANA[0]))
    for i, fila in enumerate(DIAS_SEMANA[1:], 5):
        _row(ws, i, list(fila))
    ws["A11"] = "Meses en orden calendario (01-ene a 31-dic) y numero de dias (año no bisiesto):"
    _hdr(ws, 12, ["Mes", "Dias"])
    for i, (m, d) in enumerate(MESES_CAL, 13):
        _row(ws, i, [m, d])
    ws.cell(25, 1, "TOTAL"); ws.cell(25, 2, 365)
    for c in ("A", "B", "C", "D", "E", "F"):
        ws.column_dimensions[c].width = 14

    wb.save(path)
    print("OK ->", path)


if __name__ == "__main__":
    import sys
    construir(sys.argv[1] if len(sys.argv) > 1 else "Parametros_Modelo.xlsx")
