"""
modelo.py  -- MODELO PRINCIPAL
==============================
Lee 'Parametros_Modelo.xlsx', carga los afluentes, busca la fecha optima de
inicio del mantenimiento de Ralco (min. vertimiento esperado Ralco+Pangue),
y genera las salidas (fecha optima, curva de barrido, bandas de probabilidad,
evolucion de cota, series diarias por banda).

Ejecutar:   python src/modelo.py   [ruta_Parametros.xlsx]
"""
import os, sys, time
import numpy as np
import pandas as pd
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import hidraulica as H
import datos as D
from optimizador import optimizar_escenario

FECHAS = pd.date_range("2001-01-01", "2001-12-31", freq="D")  # año no bisiesto (rotulos)


# ------------------------------------------------------------------ util
def leer_parametros(path_xlsx):
    wb = openpyxl.load_workbook(path_xlsx, data_only=True)
    ws = wb["Parametros"]
    p = {}
    for r in range(4, ws.max_row + 1):
        clave = ws.cell(r, 1).value
        if clave is None or str(clave).startswith(("---", " ")):
            continue
        val = ws.cell(r, 2).value
        if val is not None:
            p[str(clave).strip()] = val
    return p


def construir_mascaras(par):
    mpd = []
    for im, mes in enumerate(D.MESES_CAL):
        mpd += [im + 1] * D.DIAS_MES[mes]
    mpd = np.array(mpd)
    riego = np.isin(mpd, list(range(int(par["riego_mes_ini"]), int(par["riego_mes_fin"]) + 1)))
    return mpd, riego


def mask_mantenimiento(d0, dur, T, cruce=True):
    m = np.zeros(T, bool)
    if cruce:
        idx = (np.arange(d0, d0 + dur)) % T
    else:
        idx = np.arange(d0, min(d0 + dur, T))
    m[idx] = True
    return m


# ------------------------------------------------------------------ nucleo
def vertimiento_medio(R, P, par, d0, dur, T, n_iter):
    """Vertimiento esperado (MWh) promedio sobre escenarios para una fecha/duracion."""
    par = dict(par); par["n_iter"] = n_iter
    mant = mask_mantenimiento(d0, dur, T, par["_cruce"])
    tot = []
    for s in range(R.shape[0]):
        r = optimizar_escenario(R[s], P[s], par, mant)
        if r is None:
            tot.append(np.nan)
        else:
            tot.append(r["vert_total_MWh"])
    return np.nanmean(tot)


def buscar_fecha_optima(R, P, par, dur, T, paso, refina, n_iter, log=print):
    # barrido grueso
    starts = list(range(0, T, paso))
    curva = {}
    for d0 in starts:
        curva[d0] = vertimiento_medio(R, P, par, d0, dur, T, n_iter)
        log(f"    dur={dur:>3}  inicio dia {d0+1:>3} ({FECHAS[d0].strftime('%d-%b')})"
            f"  vert.esp={curva[d0]/1000:8.1f} GWh")
    best = min(curva, key=curva.get)
    # refinamiento fino alrededor del mejor
    lo, hi = best - paso, best + paso
    for d0 in range(lo, hi + 1, refina):
        dd = d0 % T
        if dd not in curva:
            curva[dd] = vertimiento_medio(R, P, par, dd, dur, T, n_iter)
    best = min(curva, key=curva.get)
    return best, curva


def resultados_por_banda(R, P, par, anios, d0, dur, T, n_bandas, criterio, n_iter):
    """Calcula el detalle por escenario en la fecha optima y agrupa en bandas."""
    par = dict(par); par["n_iter"] = n_iter
    mant = mask_mantenimiento(d0, dur, T, par["_cruce"])
    filas = []
    series = {}   # s -> dict de series diarias
    for s in range(R.shape[0]):
        r = optimizar_escenario(R[s], P[s], par, mant)
        if r is None:
            continue
        afl_anual = R[s].mean() + P[s].mean()
        filas.append(dict(esc=s, anio=anios[s], afluente_m3s=afl_anual,
                          gen_ralco=r["gen_ralco_MWh"], gen_pangue=r["gen_pangue_MWh"],
                          gen_total=r["gen_total_MWh"],
                          vert_ralco=r["vert_ralco_MWh"], vert_pangue=r["vert_pangue_MWh"],
                          vert_total=r["vert_total_MWh"]))
        series[s] = dict(cota_r=r["cota_r"], cota_p=r["cota_p"],
                         gen_r=r["ene_ralco_serie"], gen_p=r["ene_pangue_serie"],
                         vert_r=r["vert_ralco_serie"], vert_p=r["vert_pangue_serie"],
                         afl_r=R[s], afl_p=P[s])
    df = pd.DataFrame(filas)
    # ordenar de HUMEDO -> SECO
    key = "afluente_m3s" if criterio == "afluente" else "vert_total"
    df = df.sort_values(key, ascending=False).reset_index(drop=True)
    n = len(df)
    # banda 1 = 0-20% (humedo/mayor vertimiento) ... banda 5 = 80-100% (seco/menor)
    df["banda"] = (np.floor(np.arange(n) / n * n_bandas) + 1).astype(int).clip(1, n_bandas)
    etiquetas = {b: f"{int((b-1)/n_bandas*100)}-{int(b/n_bandas*100)}%" for b in range(1, n_bandas + 1)}
    df["rango_prob"] = df["banda"].map(etiquetas)
    return df, series, etiquetas


# ------------------------------------------------------------------ main
def main(path_param):
    t0 = time.time()
    base = os.path.dirname(os.path.abspath(path_param))  # carpeta raiz del paquete (donde esta el Excel)
    par_xls = leer_parametros(path_param)

    # rutas de datos (relativas a la raiz del paquete)
    def ruta(p):
        return p if os.path.isabs(p) else os.path.join(base, p)
    pr = ruta(par_xls.get("ruta_afluentes_ralco"))
    pp = ruta(par_xls.get("ruta_afluentes_pangue"))
    outdir = ruta(par_xls.get("carpeta_salidas", "salidas"))
    os.makedirs(outdir, exist_ok=True)

    print("Cargando afluentes...")
    aR, R = D.construir_series_diarias(pr)
    aP, P = D.construir_series_diarias(pp)
    assert aR == aP, "Los años de ambos archivos no coinciden"
    anios = aR
    T = R.shape[1]
    print(f"  escenarios (años): {len(anios)}   dias: {T}")

    mpd, riego = construir_mascaras(par_xls)

    par = dict(
        eco=float(par_xls["caudal_ecologico_m3s"]),
        P_ralco_max=float(par_xls["ralco_pot_max_MW"]),
        P_pangue_max=float(par_xls["pangue_pot_max_MW"]),
        P_pangue_min=float(par_xls["pangue_pot_min_MW"]),
        drop_high=float(par_xls["ralco_baja_sobre_umbral_cm"]) / 100.0,
        drop_low=float(par_xls["ralco_baja_bajo_umbral_cm"]) / 100.0,
        cota_thr=float(par_xls["ralco_cota_umbral_baja"]),
        Vr_ini=H.volumen_ralco(float(par_xls["ralco_cota_ini"])),
        Vr_min=H.volumen_ralco(float(par_xls["ralco_cota_min_op"])),
        Vr_max=H.volumen_ralco(float(par_xls["ralco_cota_max"])),
        Vp_ini=H.volumen_pangue(float(par_xls["pangue_cota_ini"])),
        Vp_min=H.volumen_pangue(float(par_xls["pangue_cota_min_op"])),
        Vp_max=H.volumen_pangue(float(par_xls["pangue_cota_max"])),
        riego_mask=riego,
        _cruce=str(par_xls.get("mant_permite_cruce_anio", "SI")).upper().startswith("S"),
    )
    duraciones = [int(x) for x in str(par_xls["mant_duraciones_dias"]).split(",")]
    paso = int(par_xls["busqueda_paso_dias"]); refina = int(par_xls["busqueda_refina_dias"])
    n_it_b = int(par_xls["n_iter_busqueda"]); n_it_f = int(par_xls["n_iter_final"])
    n_bandas = int(par_xls["n_bandas_probabilidad"])
    criterio = str(par_xls.get("criterio_bandas", "afluente"))

    # ---- barrido y fecha optima por duracion ----
    resumen = []
    barrido_rows = []
    detalle_guardado = False
    for dur in duraciones:
        print(f"\n>>> Buscando fecha optima  (mantenimiento {dur} dias)")
        best, curva = buscar_fecha_optima(R, P, par, dur, T, paso, refina, n_it_b)
        for d0 in sorted(curva):
            barrido_rows.append(dict(duracion_dias=dur, inicio_dia=d0 + 1,
                                     fecha_inicio=FECHAS[d0].strftime("%d-%b"),
                                     vert_esperado_GWh=round(curva[d0] / 1000, 2)))
        f_ini = FECHAS[best]
        f_fin = FECHAS[(best + dur - 1) % T]
        print(f"    OPTIMO: inicio dia {best+1} ({f_ini.strftime('%d-%b')}) "
              f"-> fin {f_fin.strftime('%d-%b')}   vert.esp={curva[best]/1000:.1f} GWh")

        # detalle por banda en la fecha optima
        df, series, etiquetas = resultados_por_banda(
            R, P, par, anios, best, dur, T, n_bandas, criterio, n_it_f)

        # KPIs banda
        agg = df.groupby(["banda", "rango_prob"]).agg(
            n_escenarios=("esc", "count"),
            afluente_medio_m3s=("afluente_m3s", "mean"),
            gen_total_GWh=("gen_total", lambda x: x.mean() / 1000),
            vert_ralco_GWh=("vert_ralco", lambda x: x.mean() / 1000),
            vert_pangue_GWh=("vert_pangue", lambda x: x.mean() / 1000),
            vert_total_GWh=("vert_total", lambda x: x.mean() / 1000),
        ).reset_index()
        agg.insert(0, "duracion_dias", dur)
        agg.insert(1, "fecha_inicio_opt", f_ini.strftime("%d-%b"))

        resumen.append(dict(
            duracion_dias=dur, inicio_dia_opt=best + 1,
            fecha_inicio_opt=f_ini.strftime("%d-%b"),
            fecha_fin_opt=f_fin.strftime("%d-%b"),
            vert_esperado_GWh=round(df["vert_total"].mean() / 1000, 2),  # 3-iter (consistente con bandas)
            vert_ralco_GWh=round(df["vert_ralco"].mean() / 1000, 2),
            vert_pangue_GWh=round(df["vert_pangue"].mean() / 1000, 2),
            gen_total_media_GWh=round(df["gen_total"].mean() / 1000, 1),
            vert_esperado_barrido_GWh=round(curva[best] / 1000, 2),      # 1-iter (referencia del barrido)
        ))

        # guardar bandas (append por duracion)
        modo = "w" if dur == duraciones[0] else "a"
        hdr = dur == duraciones[0]
        agg.round(2).to_csv(os.path.join(outdir, "bandas_probabilidad.csv"),
                            index=False, mode=modo, header=hdr)

        # cota + series diarias por banda SOLO para la duracion base (primera)
        if not detalle_guardado:
            _guardar_series_banda(df, series, etiquetas, dur, best, FECHAS, T, outdir, mpd)
            detalle_guardado = True

    pd.DataFrame(resumen).to_csv(os.path.join(outdir, "resumen_fecha_optima.csv"), index=False)
    df_barr = pd.DataFrame(barrido_rows)
    df_barr.to_csv(os.path.join(outdir, "barrido_fechas.csv"), index=False)

    # ---- graficos ----
    try:
        _graficar(outdir, df_barr, resumen, duraciones)
    except Exception as e:
        print("  (aviso: no se generaron graficos:", e, ")")

    print(f"\nListo en {time.time()-t0:.1f} s. Salidas en: {outdir}")
    print(pd.DataFrame(resumen).to_string(index=False))


def _guardar_series_banda(df, series, etiquetas, dur, best, FECHAS, T, outdir, mpd):
    """Promedia por banda las series diarias y guarda cota + detalle."""
    bandas = sorted(df["banda"].unique())
    fechas = [FECHAS[t].strftime("%d-%b") for t in range(T)]
    cota_cols = {"dia": np.arange(1, T + 1), "fecha": fechas}
    det_frames = []
    for b in bandas:
        escs = df[df["banda"] == b]["esc"].tolist()
        cr = np.mean([series[s]["cota_r"] for s in escs], axis=0)
        cp = np.mean([series[s]["cota_p"] for s in escs], axis=0)
        gr = np.mean([series[s]["gen_r"] for s in escs], axis=0)
        gp = np.mean([series[s]["gen_p"] for s in escs], axis=0)
        vr = np.mean([series[s]["vert_r"] for s in escs], axis=0)
        vp = np.mean([series[s]["vert_p"] for s in escs], axis=0)
        ar = np.mean([series[s]["afl_r"] for s in escs], axis=0)
        ap = np.mean([series[s]["afl_p"] for s in escs], axis=0)
        et = etiquetas[b]
        cota_cols[f"cota_Ralco_msnm_{et}"] = np.round(cr, 2)
        cota_cols[f"cota_Pangue_msnm_{et}"] = np.round(cp, 3)
        det_frames.append(pd.DataFrame(dict(
            dia=np.arange(1, T + 1), fecha=fechas, banda=b, rango_prob=et,
            afluente_ralco_m3s=np.round(ar, 1), afluente_pangue_m3s=np.round(ap, 1),
            cota_ralco_msnm=np.round(cr, 2), cota_pangue_msnm=np.round(cp, 3),
            gen_ralco_MWh=np.round(gr, 1), gen_pangue_MWh=np.round(gp, 1),
            vert_ralco_MWh=np.round(vr, 1), vert_pangue_MWh=np.round(vp, 1),
        )))
    pd.DataFrame(cota_cols).to_csv(os.path.join(outdir, "cota_evolucion_bandas.csv"), index=False)
    pd.concat(det_frames).to_csv(os.path.join(outdir, "detalle_diario_bandas.csv"), index=False)


def _graficar(outdir, df_barr, resumen, duraciones):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    # 1) curva vertimiento esperado vs fecha de inicio (por duracion)
    fig, ax = plt.subplots(figsize=(10, 5))
    for dur in duraciones:
        d = df_barr[df_barr["duracion_dias"] == dur].sort_values("inicio_dia")
        ax.plot(d["inicio_dia"], d["vert_esperado_GWh"], marker="o", ms=3, label=f"{dur} dias")
    for r in resumen:
        ax.scatter([r["inicio_dia_opt"]], [r["vert_esperado_barrido_GWh"]],
                   color="red", zorder=5, s=40)
    ax.set_xlabel("Dia de inicio del mantenimiento (1 = 01-ene)")
    ax.set_ylabel("Vertimiento esperado [GWh]")
    ax.set_title("Vertimiento esperado Ralco+Pangue vs fecha de inicio del mantenimiento")
    ax.grid(alpha=.3); ax.legend(title="Duracion")
    fig.tight_layout(); fig.savefig(os.path.join(outdir, "fig_barrido_fechas.png"), dpi=120)
    plt.close(fig)

    # 2) evolucion de cota por banda (duracion base)
    cota = pd.read_csv(os.path.join(outdir, "cota_evolucion_bandas.csv"))
    fig, ax = plt.subplots(figsize=(11, 5))
    for col in [c for c in cota.columns if c.startswith("cota_Ralco")]:
        et = col.replace("cota_Ralco_msnm_", "")
        ax.plot(cota["dia"], cota[col], label=et)
    ax.axhline(725, color="k", lw=.6, ls="--"); ax.axhline(692, color="k", lw=.6, ls="--")
    ax.axhline(708, color="gray", lw=.6, ls=":")
    ax.set_xlabel("Dia del año (1 = 01-ene)")
    ax.set_ylabel("Cota Ralco [msnm]")
    ax.set_title(f"Evolucion de cota Ralco por banda de probabilidad (mantenimiento {duraciones[0]} dias)")
    ax.grid(alpha=.3); ax.legend(title="Banda (humedo->seco)")
    fig.tight_layout(); fig.savefig(os.path.join(outdir, "fig_cota_ralco_bandas.png"), dpi=120)
    plt.close(fig)


if __name__ == "__main__":
    ruta = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "Parametros_Modelo.xlsx")
    main(ruta)
